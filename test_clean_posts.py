from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from docx import Document

from clean_posts import clean_docx


class CleanPostsTests(unittest.TestCase):
    def test_clean_docx_normalizes_quotes_and_preserves_formatting(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "post.docx"
            document = Document()
            run = document.add_paragraph().add_run("‘one’ “two”")
            run.bold = True
            document.save(path)

            self.assertEqual(clean_docx(path), 4)

            result = Document(path)
            self.assertEqual(result.paragraphs[0].text, "'one' \"two\"")
            self.assertTrue(result.paragraphs[0].runs[0].bold)


if __name__ == "__main__":
    unittest.main()
