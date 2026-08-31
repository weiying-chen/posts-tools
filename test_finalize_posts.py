from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from docx import Document

from clean_posts import clean_docx
from finalize_posts import main


class FinalizePostsTests(unittest.TestCase):
    def test_normalizes_smart_quotes_before_phrase_check(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "post.docx"
            document = Document()
            paragraph = document.add_paragraph()
            run = paragraph.add_run(
                "The Master said “Let’s protect people’s health.” Let’s take a listen."
            )
            run.bold = True
            document.add_paragraph("一起來聽聽。")
            document.add_paragraph("參考資料：")
            document.save(path)

            self.assertEqual(main([str(path)]), 0)

            result = Document(path)
            self.assertEqual(
                result.paragraphs[0].text,
                'The Master said "Let\'s protect people\'s health." Let\'s take a listen.',
            )
            self.assertTrue(result.paragraphs[0].runs[0].bold)

    def test_copy_mode_cleans_copy_without_changing_source(self) -> None:
        with TemporaryDirectory() as directory:
            source = Path(directory) / "post.docx"
            document = Document()
            document.add_paragraph("Let’s take a listen.")
            document.add_paragraph("一起來聽聽。")
            document.add_paragraph("參考資料：")
            document.save(source)

            self.assertEqual(main(["--copy", str(source)]), 0)

            cleaned = source.with_name("post_finalized.docx")
            self.assertIn("Let’s", Document(source).paragraphs[0].text)
            self.assertIn("Let's", Document(cleaned).paragraphs[0].text)

    def test_normalizer_reports_replaced_characters(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "post.docx"
            document = Document()
            document.add_paragraph("‘one’ “two”")
            document.save(path)

            self.assertEqual(clean_docx(path), 4)
            self.assertEqual(Document(path).paragraphs[0].text, "'one' \"two\"")


if __name__ == "__main__":
    unittest.main()
