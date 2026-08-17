from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from docx import Document
from docx.enum.text import WD_COLOR_INDEX

from posts_highlight import highlight_docx


class ReferenceHighlightTests(unittest.TestCase):
    def test_reference_text_becomes_cyan_without_overwriting_green(self) -> None:
        with TemporaryDirectory() as directory:
            source = Path(directory) / "source.docx"
            destination = Path(directory) / "destination.docx"
            document = Document()
            before = document.add_paragraph("post text")
            document.add_paragraph("參考資料：")
            reference = document.add_paragraph()
            reference.add_run("ordinary source text")
            wrong_color = reference.add_run("wrong-color source text")
            wrong_color.font.highlight_color = WD_COLOR_INDEX.YELLOW
            green = reference.add_run("green source text")
            green.font.highlight_color = WD_COLOR_INDEX.BRIGHT_GREEN
            document.save(source)

            result = highlight_docx(source, destination)

            self.assertEqual(result.green_paragraphs, 0)
            self.assertEqual(result.cyan_paragraphs, 1)
            self.assertEqual(result.total_paragraphs, 1)
            self.assertEqual(result.skipped_hyperlink_paragraphs, 0)
            result = Document(destination)
            self.assertIsNone(result.paragraphs[0].runs[0].font.highlight_color)
            self.assertIsNone(result.paragraphs[1].runs[0].font.highlight_color)
            self.assertEqual(
                result.paragraphs[2].runs[0].font.highlight_color,
                WD_COLOR_INDEX.TURQUOISE,
            )
            self.assertEqual(
                result.paragraphs[2].runs[1].font.highlight_color,
                WD_COLOR_INDEX.TURQUOISE,
            )
            self.assertEqual(
                result.paragraphs[2].runs[2].font.highlight_color,
                WD_COLOR_INDEX.BRIGHT_GREEN,
            )

    def test_star_marking_wins_over_automatic_reference_cyan(self) -> None:
        with TemporaryDirectory() as directory:
            source = Path(directory) / "source.docx"
            destination = Path(directory) / "destination.docx"
            document = Document()
            document.add_paragraph("參考資料:")
            document.add_paragraph("cyan *green* cyan")
            document.save(source)

            result = highlight_docx(source, destination)

            runs = Document(destination).paragraphs[1].runs
            self.assertEqual(result.green_paragraphs, 1)
            self.assertEqual(result.cyan_paragraphs, 1)
            self.assertEqual(result.total_paragraphs, 1)
            self.assertEqual([run.text for run in runs], ["cyan ", "green", " cyan"])
            self.assertEqual(
                [run.font.highlight_color for run in runs],
                [
                    WD_COLOR_INDEX.TURQUOISE,
                    WD_COLOR_INDEX.BRIGHT_GREEN,
                    WD_COLOR_INDEX.TURQUOISE,
                ],
            )

    def test_reference_line_breaks_are_not_highlighted(self) -> None:
        with TemporaryDirectory() as directory:
            source = Path(directory) / "source.docx"
            destination = Path(directory) / "destination.docx"
            document = Document()
            document.add_paragraph("參考資料：")
            paragraph = document.add_paragraph()
            blank = paragraph.add_run("\n")
            blank.font.highlight_color = WD_COLOR_INDEX.TURQUOISE
            paragraph.add_run("\nvisible text\n\nmore text")
            paragraph.add_run(" ")
            blank_paragraph = document.add_paragraph(" ")
            blank_paragraph.runs[0].font.highlight_color = WD_COLOR_INDEX.TURQUOISE
            document.save(source)

            highlight_docx(source, destination)

            runs = Document(destination).paragraphs[1].runs
            self.assertEqual(
                [run.text for run in runs],
                ["\n", "\n", "visible text", "\n\n", "more text", " "],
            )
            self.assertEqual(
                [run.font.highlight_color for run in runs],
                [
                    None,
                    None,
                    WD_COLOR_INDEX.TURQUOISE,
                    None,
                    WD_COLOR_INDEX.TURQUOISE,
                    WD_COLOR_INDEX.TURQUOISE,
                ],
            )
            self.assertIsNone(
                Document(destination).paragraphs[2].runs[0].font.highlight_color
            )


if __name__ == "__main__":
    unittest.main()
